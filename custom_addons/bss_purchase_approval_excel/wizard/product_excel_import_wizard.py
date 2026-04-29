import base64
from io import BytesIO

from odoo import _, fields, models
from odoo.exceptions import UserError


class ProductExcelImportWizard(models.TransientModel):
    _name = "product.excel.import.wizard"
    _description = "Product Excel Import Wizard"

    file_data = fields.Binary(string="Excel File", required=True)
    file_name = fields.Char(string="File Name")

    def _normalize_header(self, header):
        return str(header or "").strip().lower().replace("*", "")

    def _to_bool(self, value):
        if isinstance(value, bool):
            return value
        if value is None:
            return False
        txt = str(value).strip().lower()
        return txt in {"1", "true", "yes", "y"}

    def _to_float(self, value, field_label, row_number):
        if value in (None, ""):
            return None
        try:
            return float(value)
        except Exception as exc:
            raise UserError(_("Row %s: %s must be a number.") % (row_number, field_label)) from exc

    def _row_value(self, values, index):
        if index is None or index >= len(values):
            return None
        return values[index]

    def _get_category(self, category_name, row_number):
        if not category_name:
            return False

        category_name = str(category_name).strip()
        category = self.env["product.category"].search([("name", "=", category_name)], limit=1)
        if not category:
            category = self.env["product.category"].search(
                [("complete_name", "=", category_name)],
                limit=1,
            )
        if not category:
            raise UserError(
                _("Row %s: Category '%s' does not exist.") % (row_number, category_name)
            )
        return category

    def _get_uom(self, uom_name, row_number, field_label):
        if not uom_name:
            return False

        uom_name = str(uom_name).strip()
        uom = self.env["uom.uom"].search([("name", "=", uom_name)], limit=1)
        if not uom:
            raise UserError(
                _("Row %s: %s '%s' does not exist.")
                % (row_number, field_label, uom_name)
            )
        return uom

    def action_import_products(self):
        self.ensure_one()
        try:
            from openpyxl import load_workbook
        except Exception as exc:
            raise UserError(_("openpyxl is required to import Excel files.")) from exc

        workbook = load_workbook(
            filename=BytesIO(base64.b64decode(self.file_data)),
            data_only=True,
            read_only=True,
        )
        try:
            sheet = workbook.worksheets[0]

            rows = sheet.iter_rows(values_only=True)
            headers_row = next(rows, None)
            if not headers_row:
                raise UserError(_("The uploaded file is empty."))

            headers = [self._normalize_header(h) for h in headers_row]
            col_index = {name: idx for idx, name in enumerate(headers)}

            code_col = col_index.get("product code")
            if code_col is None:
                code_col = col_index.get("internal reference")

            required_columns = {
                "name": col_index.get("name"),
                "product type": col_index.get("product type"),
                "product code": code_col,
                "unit of measure": col_index.get("unit of measure"),
                "purchase uom": col_index.get("purchase uom"),
            }
            missing_columns = [label for label, idx in required_columns.items() if idx is None]
            if missing_columns:
                raise UserError(
                    _("Missing required columns: %s") % ", ".join(missing_columns)
                )

            allowed_types = {"consu", "service", "combo"}
            seen_codes = set()
            created_count = 0
            updated_count = 0

            for row_number, row in enumerate(rows, start=2):
                values = list(row) if row else []
                if not any(v not in (None, "") for v in values):
                    continue

                code = str(self._row_value(values, code_col) or "").strip()
                if not code:
                    raise UserError(_("Row %s: Product Code is required.") % row_number)

                if code in seen_codes:
                    raise UserError(
                        _("Row %s: Duplicate Product Code '%s' in the uploaded file.")
                        % (row_number, code)
                    )
                seen_codes.add(code)

                products = self.env["product.template"].with_context(active_test=False).search(
                    [("default_code", "=", code)]
                )
                if len(products) > 1:
                    raise UserError(
                        _("Row %s: Product Code '%s' already exists on multiple products.")
                        % (row_number, code)
                    )
                product = products[:1]

                name = str(self._row_value(values, col_index["name"]) or "").strip()
                ptype = str(self._row_value(values, col_index["product type"]) or "").strip().lower()
                category_name = str(self._row_value(values, col_index.get("category")) or "").strip()
                uom_name = str(self._row_value(values, col_index["unit of measure"]) or "").strip()
                purchase_uom_name = (
                    str(self._row_value(values, col_index["purchase uom"]) or "").strip()
                )

                if not name:
                    raise UserError(_("Row %s: Name is required.") % row_number)
                if not ptype:
                    raise UserError(_("Row %s: Product Type is required.") % row_number)
                if ptype not in allowed_types:
                    raise UserError(
                        _(
                            "Row %s: Product Type '%s' is invalid. Use one of: consu, service, combo."
                        )
                        % (row_number, ptype)
                    )
                if not uom_name:
                    raise UserError(_("Row %s: Unit of Measure is required.") % row_number)
                if not purchase_uom_name:
                    raise UserError(_("Row %s: Purchase UoM is required.") % row_number)

                vals = {
                    "default_code": code,
                    "name": name,
                    "type": ptype,
                }

                category = self._get_category(category_name, row_number)
                if category:
                    vals["categ_id"] = category.id

                uom = self._get_uom(uom_name, row_number, "Unit of Measure")
                purchase_uom = self._get_uom(purchase_uom_name, row_number, "Purchase UoM")
                vals["uom_id"] = uom.id
                vals["uom_po_id"] = purchase_uom.id

                list_price = self._to_float(
                    self._row_value(values, col_index.get("sales price")),
                    "Sales Price",
                    row_number,
                )
                if list_price is not None:
                    vals["list_price"] = list_price

                standard_price = self._to_float(
                    self._row_value(values, col_index.get("cost")),
                    "Cost",
                    row_number,
                )
                if standard_price is not None:
                    vals["standard_price"] = standard_price

                can_be_sold = self._row_value(values, col_index.get("can be sold"))
                if can_be_sold is not None:
                    vals["sale_ok"] = self._to_bool(can_be_sold)

                can_be_purchased = self._row_value(values, col_index.get("can be purchased"))
                if can_be_purchased is not None:
                    vals["purchase_ok"] = self._to_bool(can_be_purchased)

                if product:
                    product.write(vals)
                    updated_count += 1
                else:
                    self.env["product.template"].create(vals)
                    created_count += 1

            return {
                "type": "ir.actions.client",
                "tag": "display_notification",
                "params": {
                    "title": _("Import Completed"),
                    "message": _("Created: %s, Updated: %s") % (created_count, updated_count),
                    "type": "success",
                    "sticky": False,
                },
            }
        finally:
            workbook.close()
